import os
import sys
import tempfile
import unittest

PYTHON_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_DIR = os.path.join(PYTHON_DIR, "src")
if SRC_DIR not in sys.path:
    sys.path.insert(0, SRC_DIR)

from rag_project.document_processing.converter import (
    convert_to_markdown,
    _convert_txt_to_markdown,
    _convert_pdf_to_markdown,
    _convert_docx_to_markdown,
)
from docx import Document as DocxDocument
from pypdf import PdfWriter


class ConverterTest(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_convert_markdown_file(self):
        md_file = os.path.join(self.temp_dir.name, "sample.md")
        with open(md_file, "w", encoding="utf-8") as f:
            f.write("# 測試會議\n\n討論重點...")

        out_md = convert_to_markdown(md_file, output_dir=self.temp_dir.name)
        self.assertTrue(os.path.exists(out_md))
        with open(out_md, "r", encoding="utf-8") as f:
            content = f.read()
        self.assertIn("# 測試會議", content)

    def test_convert_txt_file(self):
        txt_file = os.path.join(self.temp_dir.name, "notes.txt")
        with open(txt_file, "w", encoding="utf-8") as f:
            f.write("這是純文字會議內容。")

        out_md = convert_to_markdown(txt_file, output_dir=self.temp_dir.name)
        self.assertTrue(os.path.exists(out_md))
        with open(out_md, "r", encoding="utf-8") as f:
            content = f.read()
        self.assertIn("# notes", content)
        self.assertIn("這是純文字會議內容。", content)

    def test_convert_docx_file(self):
        docx_file = os.path.join(self.temp_dir.name, "meeting.docx")
        doc = DocxDocument()
        doc.add_heading("Word 會議標題", level=1)
        doc.add_paragraph("Word 會議內容段落。")
        doc.save(docx_file)

        out_md = convert_to_markdown(docx_file, output_dir=self.temp_dir.name)
        self.assertTrue(os.path.exists(out_md))
        with open(out_md, "r", encoding="utf-8") as f:
            content = f.read()
        self.assertIn("Word 會議標題", content)
        self.assertIn("Word 會議內容段落。", content)

    def test_convert_pdf_file(self):
        pdf_file = os.path.join(self.temp_dir.name, "document.pdf")
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        with open(pdf_file, "wb") as f:
            writer.write(f)

        out_md = convert_to_markdown(pdf_file, output_dir=self.temp_dir.name)
        self.assertTrue(os.path.exists(out_md))
        self.assertTrue(out_md.endswith(".md"))


if __name__ == "__main__":
    unittest.main()
